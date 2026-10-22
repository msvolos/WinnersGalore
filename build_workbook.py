"""Build "DRF Tracker.xlsx" from ledger.csv.

ledger.csv stays the source of truth (the daily agent appends and grades rows
there). This script rebuilds the formatted workbook from it, comparing two
betting strategies on each Race of the Day:

  Top Pick:    $2 Win/Place/Show ($6) on the top pick only
  Top 2 Picks: $6 WPS on the top pick plus $6 WPS on the second choice ($12)

It also shows the second pick on its own so you can see what it adds. Every KPI
is an Excel formula over the Ledger sheet, so edits there recalculate.

Usage: python build_workbook.py [--csv ledger.csv] [--out "DRF Tracker.xlsx"]
"""
import argparse
import csv
import os
import sys
import tempfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
FONT = "Arial"
MAX_ROW = 5000  # KPI formulas cover Ledger rows 2..MAX_ROW

GREEN = "1F4E3D"   # top pick
GOLD = "C9A227"    # second pick
BLUE = "2F5597"    # top 2 combined
LIGHT = "E8F0EC"
GREY = "F2F2F2"
MONEY = '$#,##0.00;[Red]($#,##0.00);"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
ODDS = '0.0"-1"'
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# (csv field, Ledger header). Ledger columns follow this order.
CSV_FIELDS = [
    ("race_date", "Race Date"), ("track_code", "Track"), ("track_name", "Track Name"),
    ("race_number", "Race #"), ("surface", "Surface / Distance"), ("pp_url", "DRF PP Link"),
    ("pick_program", "Pick #"), ("pick_horse", "Top Pick"), ("pick_ml_odds", "Pick ML Odds"),
    ("second_choice", "Second Pick"), ("handicapped_on", "Handicapped On"), ("status", "Pick Status"),
    ("finish_position", "Pick Finish"), ("win_payout", "Pick Win Paid"),
    ("place_payout", "Pick Place Paid"), ("show_payout", "Pick Show Paid"),
    ("wagered", "Pick Wagered"), ("returned", "Pick Returned"), ("net", "Pick Net"),
    ("winner", "Race Winner"), ("winner_win_payout", "Winner $2 Win"), ("graded_on", "Graded On"),
    ("notes", "Notes"),
    ("second_program", "2nd #"), ("second_ml_odds", "2nd ML Odds"), ("second_status", "2nd Status"),
    ("second_finish", "2nd Finish"), ("second_win_payout", "2nd Win Paid"),
    ("second_place_payout", "2nd Place Paid"), ("second_show_payout", "2nd Show Paid"),
    ("second_wagered", "2nd Wagered"), ("second_returned", "2nd Returned"), ("second_net", "2nd Net"),
]
HELPERS = [
    ("p_odds", "Pick Odds (to 1)"), ("p_band", "Pick Odds Band"), ("p_fin", "Pick Finish Group"),
    ("month", "Month"), ("s_odds", "2nd Odds (to 1)"), ("s_band", "2nd Odds Band"),
    ("s_fin", "2nd Finish Group"),
    ("a_wag", "Top Pick Wagered"), ("a_ret", "Top Pick Returned"), ("a_net", "Top Pick Net"),
    ("b_wag", "Top 2 Wagered"), ("b_ret", "Top 2 Returned"), ("b_net", "Top 2 Net"),
    ("b_cash", "Top 2 Cashed"), ("b_won", "Top 2 Had Winner"),
]
COL = {f: get_column_letter(i) for i, (f, _) in enumerate(CSV_FIELDS + HELPERS, start=1)}

DATE_COLS = {"race_date", "handicapped_on", "graded_on"}
NUM_COLS = {"race_number", "finish_position", "win_payout", "place_payout", "show_payout",
            "wagered", "winner_win_payout", "second_finish", "second_win_payout",
            "second_place_payout", "second_show_payout", "second_wagered"}
FORMULA_FIELDS = {"returned", "net", "second_returned", "second_net"}
SETTLED = ("graded", "scratched", "void", "no_play")
BANDS = ["Under 7-2", "7-2 to 5-1", "6-1 to 10-1", "Over 10-1"]
FINISHES = ["1st", "2nd", "3rd", "4th or worse"]
G = '"graded"'


def font(**kw):
    kw.setdefault("name", FONT)
    kw.setdefault("size", 10)
    return Font(**kw)


def L(field):
    """Absolute Ledger range for a ledger field."""
    c = COL[field]
    return f"Ledger!${c}$2:${c}${MAX_ROW}"


def parse_value(field, raw):
    raw = (raw or "").strip()
    if raw == "":
        return None
    if field in DATE_COLS:
        try:
            return datetime.strptime(raw, "%Y-%m-%d")
        except ValueError:
            return raw
    if field in NUM_COLS:
        try:
            v = float(raw.replace("$", "").replace(",", ""))
            return int(v) if v.is_integer() else v
        except ValueError:
            return raw
    return raw


def read_ledger(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = [r for r in csv.DictReader(f) if any((v or "").strip() for v in r.values())]
    rows.sort(key=lambda r: (r.get("race_date") or "", str(r.get("race_number") or "")))
    return rows


def st(r, field):
    return (r.get(field) or "").strip()


def header_row(ws, row, labels, start_col=1, fills=None):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=label)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=(fills[i] if fills else GREEN))
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def odds_formula(cell):
    return (f'=IFERROR(VALUE(LEFT({cell},FIND("-",{cell})-1))'
            f'/VALUE(MID({cell},FIND("-",{cell})+1,9)),"")')


def band_formula(cell):
    return (f'=IF({cell}="","",IF({cell}<3.5,"{BANDS[0]}",IF({cell}<=5,"{BANDS[1]}",'
            f'IF({cell}<=10,"{BANDS[2]}","{BANDS[3]}"))))')


def finish_formula(status, finish):
    return (f'=IF({status}<>"graded","",IF({finish}=1,"1st",IF({finish}=2,"2nd",'
            f'IF({finish}=3,"3rd","4th or worse"))))')


def build_ledger(wb, rows):
    ws = wb.create_sheet("Ledger")
    header_row(ws, 1, [h for _, h in CSV_FIELDS + HELPERS])
    for i, r in enumerate(rows, start=2):
        c = {f: f"{COL[f]}{i}" for f in COL}
        for j, (field, _) in enumerate(CSV_FIELDS, start=1):
            if field not in FORMULA_FIELDS:
                ws.cell(row=i, column=j, value=parse_value(field, r.get(field)))
        settled = ",".join(f'{c["status"]}="{s}"' for s in SETTLED)
        ws[c["returned"]] = f'=IF(OR({settled}),SUM({c["win_payout"]}:{c["show_payout"]}),"")'
        ws[c["net"]] = f'=IF({c["returned"]}="","",{c["returned"]}-{c["wagered"]})'
        s_settled = ",".join(f'{c["second_status"]}="{s}"' for s in SETTLED)
        ws[c["second_returned"]] = (f'=IF(OR({s_settled}),SUM({c["second_win_payout"]}:'
                                    f'{c["second_show_payout"]}),"")')
        ws[c["second_net"]] = (f'=IF({c["second_returned"]}="","",'
                               f'{c["second_returned"]}-{c["second_wagered"]})')
        ws[c["p_odds"]] = odds_formula(c["pick_ml_odds"])
        ws[c["p_band"]] = band_formula(c["p_odds"])
        ws[c["p_fin"]] = finish_formula(c["status"], c["finish_position"])
        ws[c["month"]] = f'=IF({c["race_date"]}="","",TEXT({c["race_date"]},"yyyy-mm"))'
        ws[c["s_odds"]] = odds_formula(c["second_ml_odds"])
        ws[c["s_band"]] = band_formula(c["s_odds"])
        ws[c["s_fin"]] = finish_formula(c["second_status"], c["second_finish"])
        # Strategy money per race. A bet only counts once it is graded.
        a_g, s_g = f'{c["status"]}="graded"', f'{c["second_status"]}="graded"'
        ws[c["a_wag"]] = f'=IF({a_g},{c["wagered"]},0)'
        ws[c["a_ret"]] = f'=IF({a_g},{c["returned"]},0)'
        ws[c["a_net"]] = f'={c["a_ret"]}-{c["a_wag"]}'
        ws[c["b_wag"]] = f'={c["a_wag"]}+IF({s_g},{c["second_wagered"]},0)'
        ws[c["b_ret"]] = f'={c["a_ret"]}+IF({s_g},{c["second_returned"]},0)'
        ws[c["b_net"]] = f'={c["b_ret"]}-{c["b_wag"]}'
        ws[c["b_cash"]] = f'=IF({c["b_wag"]}>0,IF({c["b_ret"]}>0,1,0),"")'
        ws[c["b_won"]] = (f'=IF({c["b_wag"]}>0,IF(OR(AND({a_g},{c["finish_position"]}=1),'
                          f'AND({s_g},{c["second_finish"]}=1)),1,0),"")')
        if r.get("pp_url"):
            ws[c["pp_url"]].hyperlink = r["pp_url"]
            ws[c["pp_url"]].value = "PPs"

    last = max(len(rows) + 1, 2)
    money = {"win_payout", "place_payout", "show_payout", "wagered", "returned", "net",
             "winner_win_payout", "second_win_payout", "second_place_payout", "second_show_payout",
             "second_wagered", "second_returned", "second_net",
             "a_wag", "a_ret", "a_net", "b_wag", "b_ret", "b_net"}
    for row in ws.iter_rows(min_row=2, max_row=last, max_col=len(COL)):
        for cell in row:
            cell.font = font()
            cell.border = BOX
            cell.alignment = Alignment(vertical="top")
    for field, col in COL.items():
        cells = ws[col][1:last]
        for cell in cells:
            if field in DATE_COLS:
                cell.number_format = "mm/dd/yyyy"
            elif field in money:
                cell.number_format = MONEY
            elif field in ("p_odds", "s_odds"):
                cell.number_format = ODDS
            elif field == "pp_url" and cell.value:
                cell.font = font(color="0563C1", underline="single")
    for field in ("net", "second_net", "a_net", "b_net"):
        rng = f"{COL[field]}2:{COL[field]}{MAX_ROW}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"],
                                      fill=PatternFill("solid", fgColor="C6EFCE")))
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"],
                                      fill=PatternFill("solid", fgColor="FFC7CE")))
    # Tint the second-pick block so the two picks are easy to tell apart.
    for field, _ in CSV_FIELDS[23:]:
        ws[f"{COL[field]}1"].fill = PatternFill("solid", fgColor="8C7319")
    for field in ("b_wag", "b_ret", "b_net", "b_cash", "b_won"):
        ws[f"{COL[field]}1"].fill = PatternFill("solid", fgColor=BLUE)
    widths = {"race_date": 11, "track_name": 15, "surface": 16, "pick_horse": 18,
              "second_choice": 18, "winner": 17, "notes": 50}
    for field, col in COL.items():
        ws.column_dimensions[col].width = widths.get(field, 10)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = f"{COL['pick_ml_odds']}2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COL))}{last}"
    return ws


def build_bet_log(wb, rows):
    """One row per race with a graded bet, in date order, plus a zero 'Start'
    row so the running charts always begin at $0."""
    ws = wb.create_sheet("Bet Log")
    cols = ["Race #", "Race Date", "Race",
            "Top Pick", "Odds", "Finish", "Wagered", "Returned", "Net",
            "Cum. Net", "Cum. ROI", "Losing Streak",
            "2nd Pick", "Odds", "Finish",
            "Top 2 Wagered", "Top 2 Returned", "Top 2 Net",
            "Top 2 Cum. Net", "Top 2 Cum. ROI", "Top 2 Losing Streak",
            "Top Pick Cum. Wagered", "Top 2 Cum. Wagered", "2nd Pick Cum. Net"]
    fills = [GREEN] * 3 + [GREEN] * 9 + [GOLD] * 3 + [BLUE] * 6 + ["808080"] * 3
    header_row(ws, 1, cols, fills=fills)
    ws["A2"], ws["C2"] = 0, "Start"
    for col in "GHIJKLPQRSTUVWX":
        ws[f"{col}2"] = 0
    out = 3
    for idx, r in enumerate(rows, start=2):
        if st(r, "status") != "graded" and st(r, "second_status") != "graded":
            continue
        n, p = out, out - 1
        lg = lambda f: f"=Ledger!{COL[f]}{idx}"  # noqa: E731
        ws[f"A{n}"] = f"=A{p}+1"
        ws[f"B{n}"] = lg("race_date")
        ws[f"C{n}"] = f'=Ledger!{COL["track_code"]}{idx}&" R"&Ledger!{COL["race_number"]}{idx}'
        ws[f"D{n}"] = lg("pick_horse")
        ws[f"E{n}"] = lg("pick_ml_odds")
        ws[f"F{n}"] = f'=IF(Ledger!{COL["status"]}{idx}="graded",Ledger!{COL["finish_position"]}{idx},Ledger!{COL["status"]}{idx})'
        ws[f"G{n}"] = lg("a_wag")
        ws[f"H{n}"] = lg("a_ret")
        ws[f"I{n}"] = lg("a_net")
        ws[f"J{n}"] = f"=J{p}+I{n}"
        ws[f"K{n}"] = f"=IFERROR(J{n}/V{n},0)"
        ws[f"L{n}"] = f"=IF(G{n}=0,L{p},IF(I{n}<0,L{p}+1,0))"
        ws[f"M{n}"] = lg("second_choice")
        ws[f"N{n}"] = lg("second_ml_odds")
        ws[f"O{n}"] = f'=IF(Ledger!{COL["second_status"]}{idx}="graded",Ledger!{COL["second_finish"]}{idx},Ledger!{COL["second_status"]}{idx})'
        ws[f"P{n}"] = lg("b_wag")
        ws[f"Q{n}"] = lg("b_ret")
        ws[f"R{n}"] = lg("b_net")
        ws[f"S{n}"] = f"=S{p}+R{n}"
        ws[f"T{n}"] = f"=IFERROR(S{n}/W{n},0)"
        ws[f"U{n}"] = f"=IF(P{n}=0,U{p},IF(R{n}<0,U{p}+1,0))"
        ws[f"V{n}"] = f"=V{p}+G{n}"
        ws[f"W{n}"] = f"=W{p}+P{n}"
        ws[f"X{n}"] = f"=S{n}-J{n}"
        out += 1
    last = out - 1
    for row in ws.iter_rows(min_row=2, max_row=last, max_col=len(cols)):
        for c in row:
            c.font = font()
            c.border = BOX
    for col in "BGHIJPQRSVWX":
        for c in ws[col][1:last]:
            c.number_format = "mm/dd/yyyy" if col == "B" else MONEY
    for col in "KT":
        for c in ws[col][1:last]:
            c.number_format = PCT
    widths = [7, 11, 9, 18, 6, 7, 9, 9, 9, 10, 9, 8, 18, 6, 7, 9, 9, 9, 10, 9, 8, 10, 10, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "D3"
    ws["Z1"] = "Built automatically from the Ledger sheet. Edit ledger.csv, not this sheet."
    ws["Z1"].font = font(italic=True, color="808080")
    return ws, last


def kpi(ws, cell, label, formula, fmt, note=None, color=GREEN):
    col, row = cell[0], int(cell[1:])
    top = ws[f"{col}{row}"]
    top.value = label
    top.font = font(size=9, bold=True, color="595959")
    top.fill = PatternFill("solid", fgColor=LIGHT)
    top.alignment = Alignment(horizontal="center", wrap_text=True)
    val = ws[f"{col}{row+1}"]
    val.value = formula
    val.number_format = fmt
    val.font = font(size=16, bold=True, color=color)
    val.fill = PatternFill("solid", fgColor=LIGHT)
    val.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 26
    ws.row_dimensions[row + 1].height = 28
    if note:
        n = ws[f"{col}{row+2}"]
        n.value = note
        n.font = font(size=8, italic=True, color="808080")
        n.alignment = Alignment(horizontal="center", wrap_text=True)


def table(ws, top, left, headers, body, fmts, fills=None):
    header_row(ws, top, headers, start_col=left, fills=fills)
    for i, values in enumerate(body, start=1):
        for j, v in enumerate(values):
            c = ws.cell(row=top + i, column=left + j, value=v)
            c.font = font(bold=j == 0)
            c.border = BOX
            c.number_format = fmts[j]
            if i % 2 == 0:
                c.fill = PatternFill("solid", fgColor=GREY)


def section(ws, cell, text):
    ws[cell] = text
    ws[cell].font = font(size=12, bold=True, color=GREEN)


def style_chart(ch, title, y_fmt=None, colors=(GREEN, GOLD, BLUE)):
    ch.title = title
    ch.style = 2
    ch.height, ch.width = 7.5, 15
    ch.legend.position = "r"
    ch.legend.overlay = False
    if y_fmt:
        ch.y_axis.numFmt = y_fmt
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.x_axis.tickLblPos = "low"
    for s, color in zip(ch.series, colors):
        if isinstance(ch, LineChart):
            s.smooth = False
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = 28000
        else:
            s.invertIfNegative = False
            s.graphicalProperties.solidFill = color


def build_dashboard(wb, rows, log_last):
    ws = wb.create_sheet("Dashboard", 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.column_dimensions["A"].width = 2
    for col in range(2, 16):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.column_dimensions["B"].width = 33

    ws["B1"] = "DRF Race of the Day: Win/Place/Show Tracker"
    ws["B1"].font = font(size=18, bold=True, color=GREEN)
    ws["B2"] = ("Top Pick = $2 Win + $2 Place + $2 Show ($6) on the top pick. Top 2 Picks = the same $6 "
                "on the top pick and on the second choice ($12 per race). Scratches are refunded. "
                "All figures are formulas over the Ledger sheet.")
    ws["B2"].font = font(size=9, italic=True, color="595959")
    ws["B3"] = '="Last updated: "&TEXT(NOW(),"mmm d, yyyy h:mm AM/PM")'
    ws["B3"].font = font(size=9, color="808080")

    races_b = f'COUNTIF({L("b_wag")},">0")'
    a_wag, a_ret = f'SUM({L("a_wag")})', f'SUM({L("a_ret")})'
    b_wag, b_ret = f'SUM({L("b_wag")})', f'SUM({L("b_ret")})'

    # KPI tiles: the head-to-head at a glance.
    kpi(ws, "B5", "TOP PICK: NET", f"={a_ret}-{a_wag}", MONEY, "$6 per race")
    kpi(ws, "D5", "TOP PICK: ROI", f"=IFERROR(({a_ret}-{a_wag})/{a_wag},0)", PCT, "net ÷ wagered")
    kpi(ws, "F5", "TOP 2 PICKS: NET", f"={b_ret}-{b_wag}", MONEY, "$12 per race", BLUE)
    kpi(ws, "H5", "TOP 2 PICKS: ROI", f"=IFERROR(({b_ret}-{b_wag})/{b_wag},0)", PCT, "net ÷ wagered", BLUE)
    kpi(ws, "J5", "2ND PICK ADDED", f"=({b_ret}-{b_wag})-({a_ret}-{a_wag})", MONEY,
        "Top 2 net minus Top Pick net", GOLD)
    kpi(ws, "L5", "RACES GRADED", f"={races_b}", "0")
    kpi(ws, "N5", "PENDING", f'=COUNTIF({L("status")},"pending")', "0", "awaiting results")

    # Strategy comparison: the main table.
    top = 10
    section(ws, f"B{top-1}", "Strategy Comparison")
    blog = wb["Bet Log"]
    last = max(log_last, 2)

    def pick_stats(status, finish, wag, ret, net_col):
        bets = f"COUNTIF({L(status)},{G})"
        return [
            f"={bets}",
            f"=SUMIFS({L(wag)},{L(status)},{G})",
            f"=SUMIFS({L(ret)},{L(status)},{G})",
            None,  # net, filled below
            None,  # ROI
            f"=IFERROR(SUMIFS({L(ret)},{L(status)},{G})/{bets},0)",
            f"=IFERROR(COUNTIFS({L(status)},{G},{L(finish)},1)/{bets},0)",
            f'=IFERROR(COUNTIFS({L(status)},{G},{L(finish)},"<=3")/{bets},0)',
            f'=IFERROR(COUNTIFS({L(status)},{G},{L(net_col)},">0")/{bets},0)',
            f'=IFERROR(MAX({L(ret)}),0)',
        ]

    a = pick_stats("status", "finish_position", "wagered", "returned", "net")
    s = pick_stats("second_status", "second_finish", "second_wagered", "second_returned", "second_net")
    b = [
        f"={races_b}", f"={b_wag}", f"={b_ret}", None, None,
        f"=IFERROR({b_ret}/{races_b},0)",
        f"=IFERROR(SUM({L('b_won')})/{races_b},0)",
        f"=IFERROR(SUM({L('b_cash')})/{races_b},0)",
        f'=IFERROR(COUNTIF({L("b_net")},">0")/{races_b},0)',
        f"=MAX({L('b_ret')})",
    ]
    labels = ["Races bet", "Total wagered", "Total returned", "Net profit / loss", "ROI",
              "Average return per race", "Win rate (a pick won)", "In-the-money rate (a ticket cashed)",
              "Races that made a profit", "Biggest single-race return",
              "Current losing streak (races)", "Longest losing streak (races)"]
    fmts_row = ["0", MONEY, MONEY, MONEY, PCT, MONEY, PCT, PCT, PCT, MONEY, "0", "0"]
    body = []
    for i, lab in enumerate(labels):
        r = top + 1 + i
        if i == 3:
            vals = [f"=C{r-1}-C{r-2}", f"=D{r-1}-D{r-2}", f"=E{r-1}-E{r-2}"]
        elif i == 4:
            vals = [f"=IFERROR(C{r-1}/C{r-3},0)", f"=IFERROR(D{r-1}/D{r-3},0)",
                    f"=IFERROR(E{r-1}/E{r-3},0)"]
        elif i == 10:
            vals = [f"=INDEX('Bet Log'!L2:L{last},{last-1})", "", f"=INDEX('Bet Log'!U2:U{last},{last-1})"]
        elif i == 11:
            vals = [f"=MAX('Bet Log'!L2:L{last})", "", f"=MAX('Bet Log'!U2:U{last})"]
        else:
            vals = [a[i], s[i], b[i]]
        body.append([lab] + vals)
    table(ws, top, 2, ["Metric", "Top Pick Only", "2nd Pick Alone", "Top 2 Picks"], body,
          ["@"] * 4, fills=[GREEN, GREEN, GOLD, BLUE])
    for i, f in enumerate(fmts_row):
        for col in "CDE":
            ws[f"{col}{top+1+i}"].number_format = f
    ws[f"B{top+len(labels)+1}"] = ("Top 2 win rate = races where either pick won; in-the-money = races "
                                    "where any of the 6 tickets paid.")
    ws[f"B{top+len(labels)+1}"].font = font(size=8, italic=True, color="808080")
    cmp_last = top + len(labels)

    # Return by bet type, per pick.
    wtop = top
    section(ws, f"G{wtop-1}", "Return by Bet Type ($2 each)")
    leg_rows = []
    for k, (name, a_col, s_col, cond) in enumerate((
            ("Win", "win_payout", "second_win_payout", "1"),
            ("Place", "place_payout", "second_place_payout", '"<=2"'),
            ("Show", "show_payout", "second_show_payout", '"<=3"'))):
        rr = wtop + 1 + k
        leg_rows.append([
            name,
            f"=2*COUNTIF({L('status')},{G})",
            f"=SUMIFS({L(a_col)},{L('status')},{G})",
            f"=IFERROR((I{rr}-H{rr})/H{rr},0)",
            f"=2*COUNTIF({L('second_status')},{G})",
            f"=SUMIFS({L(s_col)},{L('second_status')},{G})",
            f"=IFERROR((L{rr}-K{rr})/K{rr},0)",
            f"=IFERROR((I{rr}+L{rr}-H{rr}-K{rr})/(H{rr}+K{rr}),0)",
        ])
    table(ws, wtop, 7, ["Bet Type", "Top Pick Wagered", "Top Pick Returned", "Top Pick ROI",
                        "2nd Wagered", "2nd Returned", "2nd ROI", "Top 2 ROI"],
          leg_rows, ["@", MONEY, MONEY, PCT, MONEY, MONEY, PCT, PCT],
          fills=[GREEN, GREEN, GREEN, GREEN, GOLD, GOLD, GOLD, BLUE])

    # Finish distribution, per pick.
    ftop = wtop + 6
    section(ws, f"G{ftop-1}", "Where Each Pick Finished")
    fin_rows = []
    for k, lab in enumerate(FINISHES):
        rr = ftop + 1 + k
        fin_rows.append([lab, f'=COUNTIF({L("p_fin")},G{rr})', f'=COUNTIF({L("s_fin")},G{rr})'])
    table(ws, ftop, 7, ["Finish", "Top Pick", "2nd Pick"], fin_rows, ["@", "0", "0"],
          fills=[GREEN, GREEN, GOLD])

    # Odds bands, per pick.
    btop = cmp_last + 4
    section(ws, f"B{btop-1}", "Performance by Morning-Line Odds")
    band_rows = []
    for k, band in enumerate(BANDS):
        rr = btop + 1 + k
        ca = f'{L("status")},{G},{L("p_band")},B{rr}'
        cs = f'{L("second_status")},{G},{L("s_band")},B{rr}'
        band_rows.append([
            band, f"=COUNTIFS({ca})", f"=COUNTIFS({ca},{L('finish_position')},1)",
            f"=SUMIFS({L('net')},{ca})", f"=IFERROR(E{rr}/SUMIFS({L('wagered')},{ca}),0)",
            f"=COUNTIFS({cs})", f"=COUNTIFS({cs},{L('second_finish')},1)",
            f"=SUMIFS({L('second_net')},{cs})",
            f"=IFERROR(I{rr}/SUMIFS({L('second_wagered')},{cs}),0)",
        ])
    table(ws, btop, 2, ["Odds Band", "Top Pick Bets", "Wins", "Net", "ROI",
                        "2nd Pick Bets", "Wins", "Net", "ROI"],
          band_rows, ["@", "0", "0", MONEY, PCT, "0", "0", MONEY, PCT],
          fills=[GREEN] * 5 + [GOLD] * 4)

    # By track and by month (labels from the data; numbers are formulas).
    def group_table(top_row, left, title, field, labels_):
        section(ws, f"{get_column_letter(left)}{top_row-1}", title)
        out = []
        for k, lab in enumerate(labels_ or ["(none yet)"]):
            rr = top_row + 1 + k
            key = f"{get_column_letter(left)}{rr}"
            c1, c2, c3 = (get_column_letter(left + n) for n in (2, 4, 1))
            out.append([
                lab, f'=COUNTIFS({L(field)},{key},{L("b_wag")},">0")',
                f"=SUMIFS({L('a_net')},{L(field)},{key})",
                f"=IFERROR({c1}{rr}/SUMIFS({L('a_wag')},{L(field)},{key}),0)",
                f"=SUMIFS({L('b_net')},{L(field)},{key})",
                f"=IFERROR({c2}{rr}/SUMIFS({L('b_wag')},{L(field)},{key}),0)",
            ])
        table(ws, top_row, left, [title.split()[-1], "Races", "Top Pick Net", "Top Pick ROI",
                                  "Top 2 Net", "Top 2 ROI"],
              out, ["@", "0", MONEY, PCT, MONEY, PCT],
              fills=[GREEN, GREEN, GREEN, GREEN, BLUE, BLUE])
        return len(out)

    graded = [r for r in rows if "graded" in (st(r, "status"), st(r, "second_status"))]
    ttop = btop + 7
    tracks = sorted({st(r, "track_code") for r in graded} - {""})
    months = sorted({st(r, "race_date")[:7] for r in graded} - {""})
    n_t = group_table(ttop, 2, "Performance by Track", "track_code", tracks)
    n_m = group_table(ttop, 9, "Performance by Month", "month", months)

    # Charts.
    crow = ttop + max(n_t, n_m) + 3
    cats = Reference(blog, min_col=1, min_row=2, max_row=last)

    ch1 = LineChart()
    ch1.add_data(Reference(blog, min_col=10, min_row=1, max_row=last), titles_from_data=True)
    ch1.add_data(Reference(blog, min_col=19, min_row=1, max_row=last), titles_from_data=True)
    ch1.set_categories(cats)
    style_chart(ch1, "Cumulative Net: Top Pick vs Top 2 Picks (by race #)", '"$"#,##0',
                colors=(GREEN, BLUE))
    _name_series(ch1, ["Top Pick", "Top 2 Picks"])
    ws.add_chart(ch1, f"B{crow}")

    ch2 = LineChart()
    ch2.add_data(Reference(blog, min_col=11, min_row=1, max_row=last), titles_from_data=True)
    ch2.add_data(Reference(blog, min_col=20, min_row=1, max_row=last), titles_from_data=True)
    ch2.set_categories(cats)
    style_chart(ch2, "ROI Over Time (by race #)", "0%", colors=(GREEN, BLUE))
    _name_series(ch2, ["Top Pick", "Top 2 Picks"])
    ws.add_chart(ch2, f"I{crow}")

    ch3 = BarChart()
    ch3.type = "col"
    ch3.add_data(Reference(ws, min_col=10, min_row=wtop, max_row=wtop + 3), titles_from_data=True)
    ch3.add_data(Reference(ws, min_col=13, min_row=wtop, max_row=wtop + 3), titles_from_data=True)
    ch3.add_data(Reference(ws, min_col=14, min_row=wtop, max_row=wtop + 3), titles_from_data=True)
    ch3.set_categories(Reference(ws, min_col=7, min_row=wtop + 1, max_row=wtop + 3))
    style_chart(ch3, "ROI by Bet Type", "0%")
    ws.add_chart(ch3, f"B{crow+16}")

    ch4 = BarChart()
    ch4.type = "col"
    ch4.add_data(Reference(ws, min_col=8, max_col=9, min_row=ftop, max_row=ftop + 4), titles_from_data=True)
    ch4.set_categories(Reference(ws, min_col=7, min_row=ftop + 1, max_row=ftop + 4))
    style_chart(ch4, "Where Each Pick Finished (races)", "0")
    ch4.y_axis.majorUnit = 1
    ws.add_chart(ch4, f"I{crow+16}")

    ch5 = BarChart()
    ch5.type = "col"
    ch5.add_data(Reference(ws, min_col=6, min_row=btop, max_row=btop + 4), titles_from_data=True)
    ch5.add_data(Reference(ws, min_col=10, min_row=btop, max_row=btop + 4), titles_from_data=True)
    ch5.set_categories(Reference(ws, min_col=2, min_row=btop + 1, max_row=btop + 4))
    style_chart(ch5, "ROI by Morning-Line Odds", "0%")
    _name_series(ch5, ["Top Pick", "2nd Pick"])
    ws.add_chart(ch5, f"B{crow+32}")

    ch6 = BarChart()
    ch6.type = "col"
    ch6.add_data(Reference(ws, min_col=11, min_row=ttop, max_row=ttop + n_m), titles_from_data=True)
    ch6.add_data(Reference(ws, min_col=13, min_row=ttop, max_row=ttop + n_m), titles_from_data=True)
    ch6.set_categories(Reference(ws, min_col=9, min_row=ttop + 1, max_row=ttop + n_m))
    style_chart(ch6, "Net by Month", '"$"#,##0', colors=(GREEN, BLUE))
    _name_series(ch6, ["Top Pick", "Top 2 Picks"])
    ws.add_chart(ch6, f"I{crow+32}")
    return ws


def _name_series(chart, names):
    from openpyxl.chart.series import SeriesLabel
    for s, name in zip(chart.series, names):
        s.tx = SeriesLabel(v=name)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", default=os.path.join(HERE, "ledger.csv"))
    ap.add_argument("--out", default=os.path.join(HERE, "DRF Tracker.xlsx"))
    args = ap.parse_args()

    rows = read_ledger(args.csv)
    wb = Workbook()
    wb.remove(wb.active)
    build_ledger(wb, rows)
    _, log_last = build_bet_log(wb, rows)
    build_dashboard(wb, rows, log_last)
    wb.calculation.fullCalcOnLoad = True

    # Save to a temp file first so a failed save never leaves a broken workbook.
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(os.path.abspath(args.out)))
    os.close(fd)
    wb.save(tmp)
    try:
        os.replace(tmp, args.out)
    except PermissionError:
        os.remove(tmp)
        sys.exit(f"ERROR: could not overwrite {args.out} (is it open in Excel?). ledger.csv is unchanged.")
    graded = sum(1 for r in rows if st(r, "status") == "graded")
    print(f"Wrote {args.out}: {len(rows)} ledger rows, {graded} graded races.")


if __name__ == "__main__":
    main()
